"""用 data/db.json 產生新版試算表內容：sheet.json（測試用）與 .xlsx（上傳到 Google 雲端硬碟）。
用法：python3 backend/tools/build_sheet.py data/db.json 輸出資料夾 管理者Email"""
import json, subprocess, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
here = os.path.dirname(os.path.abspath(__file__))
S = json.loads(subprocess.check_output(["node", os.path.join(here, "schema.js")]))
db = json.load(open(sys.argv[1])); out = sys.argv[2]; admin = sys.argv[3]
os.makedirs(out, exist_ok=True)

def enc(k, v):
    if v is None: return ""
    if k in S["JSON_KEYS"] or isinstance(v, (list, dict)): return json.dumps(v, ensure_ascii=False, separators=(",", ":"))
    if isinstance(v, bool): return "TRUE" if v else ""
    return str(v)

tabs = {}
labels = lambda t: {k: l for k, l in S["SCHEMA"][t]}
for t, meta in S["TABLES"].items():
    keys = [k for k, _ in S["SCHEMA"][t]]
    for r in db.get(t, []):
        for k in r:
            if k not in keys: keys.append(k)
    lb = labels(t)
    head = [f"{lb.get(k, k)} ({k})" for k in keys]
    rows = [[enc(k, r.get(k)) for k in keys] for r in db.get(t, [])]
    tabs[meta["tab"]] = [head] + rows
# 合約：中控台暫不顯示，資料先保留在試算表
ck = list(dict.fromkeys(k for r in db.get("contracts", []) for k in r))
tabs["合約（暫不顯示）"] = [ck] + [[enc(k, r.get(k)) for k in ck] for r in db.get("contracts", [])]
tabs[S["USERS_TAB"]] = [["Email", "姓名", "權限（管理／編輯／檢視）", "啟用（是／否）", "最後登入"], [admin, "葉泰宏", "管理", "是", ""]]
keep = ["country", "assignStatus", "taskStatus", "milestone", "issueStatus", "issueSeverity", "issueCategory", "projectStatus"]
tabs[S["LISTS_TAB"]] = [["類別 (key)", "選項（用頓號、分隔）"]] + [[k, "、".join(db["lists"][k])] for k in keep if k in db.get("lists", {})]
tabs[S["LOG_TAB"]] = [["時間", "Email", "姓名", "分頁", "鍵", "動作", "變更內容"]]
json.dump(tabs, open(os.path.join(out, "sheet.json"), "w"), ensure_ascii=False)

wb = Workbook(); wb.remove(wb.active)
head_fill = PatternFill("solid", fgColor="E9F5B8")
for name, rows in tabs.items():
    ws = wb.create_sheet(name)
    for r in rows: ws.append(r)
    for row in ws.iter_rows():
        for c in row: c.number_format = "@"
    for c in ws[1]: c.font = Font(bold=True); c.fill = head_fill
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value or "")) for c in list(col)[:60]), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(8, w * 1.6), 48)
wb.save(os.path.join(out, "工程處中控台資料.xlsx"))
print({k: len(v) - 1 for k, v in tabs.items()})
